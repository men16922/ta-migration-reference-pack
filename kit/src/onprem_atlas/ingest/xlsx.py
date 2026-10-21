"""xlsx → 셀 문자열 표. openpyxl read_only 스트리밍이라 수천~수만 행도 메모리에 셀 객체를 쌓지 않는다."""

from __future__ import annotations

import datetime as dt
import hashlib
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


@dataclass
class Sheet:
    name: str
    rows: list[list[str]]


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Y" if value else "N"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    return " ".join(str(value).split())


def sheet_names(path: str | Path) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def read_sheet(path: str | Path, sheet: str | None = None) -> Sheet:
    """`sheet`이 없으면 첫 시트. 수식은 저장된 계산값(data_only)으로 읽는다."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet is not None and sheet not in wb.sheetnames:
            raise ValueError(f"시트 '{sheet}'가 없습니다. 시트 목록: {', '.join(wb.sheetnames)}")
        ws = wb[sheet] if sheet else wb.worksheets[0]
        if (ws.max_row or 0) > 50000 or (ws.max_column or 0) > 256:
            raise ValueError("Excel은 50,000행·256열 이하여야 합니다.")
        rows = []
        for n, row in enumerate(ws.iter_rows(values_only=True)):
            if n >= 50000 or len(row) > 256:
                raise ValueError("Excel은 50,000행·256열 이하여야 합니다.")
            rows.append([cell_text(v) for v in row])
        width = max((len(r) for r in rows), default=0)
        return Sheet(ws.title, [r + [""] * (width - len(r)) for r in rows])
    finally:
        wb.close()


def file_sha256(path: str | Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()
